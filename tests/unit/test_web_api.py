"""Career OS web dashboard API tests — FastAPI `TestClient` against
`job_agent.web.app.create_app()`. Same tmp_path/monkeypatch conventions as
the CLI's own tests (`tests/unit/test_cli_browser_approve_fill.py` etc.):
a copied config dir, a synthetic candidate dir, and a fresh sqlite file
per test, so nothing here ever touches the repo's real `data/job_agent.db`
or real `candidate/` files.
"""

from __future__ import annotations

import shutil
from datetime import UTC
from pathlib import Path

from fastapi.testclient import TestClient
from sqlalchemy import select

from job_agent.candidate.schema import (
    CandidateProfile,
    Fact,
    LocationPreferences,
    SalaryPreferences,
    TargetRoles,
    VisaInformation,
    WorkPreferences,
)
from job_agent.db.models import Job as JobRow
from job_agent.db.models import JobMatch
from job_agent.db.models import JobSource as JobSourceRow
from job_agent.db.session import get_engine, get_session_factory, init_db
from job_agent.jobs.notifications import generate_notifications
from job_agent.web.app import create_app

SYNTHETIC_VALUES = {
    "full_name": "Test Candidate",
    "email": "test.candidate@example.invalid",
    "phone": "+1-555-0100",
    "current_location": "Test City, Test Country",
    "linkedin_url": "https://linkedin.com/in/testcandidate",
}


def _fact(value: str) -> Fact[str]:
    return Fact[str](value=value, source="test_synthetic_profile", confidence=1.0, verified=True)


def _synthetic_identity_profile() -> CandidateProfile:
    unknown_pref = Fact.unknown(source="test_synthetic_profile")
    return CandidateProfile(
        identity_name=_fact(SYNTHETIC_VALUES["full_name"]),
        identity_current_location=_fact(SYNTHETIC_VALUES["current_location"]),
        contact_email=_fact(SYNTHETIC_VALUES["email"]),
        contact_phone=_fact(SYNTHETIC_VALUES["phone"]),
        contact_linkedin=_fact(SYNTHETIC_VALUES["linkedin_url"]),
        target_roles=TargetRoles(primary=("Business Analyst",)),
        work_preferences=WorkPreferences(
            remote=unknown_pref, willing_to_relocate=unknown_pref, notice_period=unknown_pref
        ),
        location_preferences=LocationPreferences(
            current_location=_fact(SYNTHETIC_VALUES["current_location"]),
            open_to_countries=unknown_pref,
        ),
        salary_preferences=SalaryPreferences(
            currency=unknown_pref,
            minimum_annual=unknown_pref,
            target_annual=unknown_pref,
            negotiable=unknown_pref,
        ),
        visa_information=VisaInformation(
            nationality=unknown_pref,
            requires_sponsorship_us=unknown_pref,
            requires_sponsorship_uk=unknown_pref,
            requires_sponsorship_eu=unknown_pref,
            requires_sponsorship_other=unknown_pref,
        ),
    )


def _make_synthetic_candidate_dir(candidate_dir: Path) -> Path:
    candidate_dir.mkdir(parents=True, exist_ok=True)
    (candidate_dir / "profile.md").write_text(
        "## Identity\n"
        f"name: {SYNTHETIC_VALUES['full_name']}\n"
        f"current_location: {SYNTHETIC_VALUES['current_location']}\n"
        "\n## Contact\n"
        f"email: {SYNTHETIC_VALUES['email']}\n"
        f"phone: {SYNTHETIC_VALUES['phone']}\n"
        f"linkedin: {SYNTHETIC_VALUES['linkedin_url']}\n"
    )
    for filename in (
        "experience.md",
        "projects.md",
        "skills.md",
        "education.md",
        "achievements.md",
    ):
        (candidate_dir / filename).write_text("")
    (candidate_dir / "answers").mkdir(exist_ok=True)

    import docx

    document = docx.Document()
    document.add_paragraph("Synthetic test resume for web API tests. No real content.")
    document.save(str(candidate_dir / "resume_master.docx"))
    return candidate_dir


def _configure_env(tmp_path: Path, monkeypatch, real_config) -> Path:
    cfg_dir = tmp_path / "config"
    shutil.copytree(real_config.env.config_dir, cfg_dir)
    candidate_dir = _make_synthetic_candidate_dir(tmp_path / "candidate")
    db_path = tmp_path / "web_api_test.db"
    monkeypatch.setenv("CONFIG_DIR", str(cfg_dir))
    monkeypatch.setenv("CANDIDATE_DIR", str(candidate_dir))
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{db_path}")
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    return db_path


def _seed_job(
    db_path: Path,
    *,
    fingerprint: str,
    source_name: str | None = None,
    lifecycle_status: str = "ACTIVE",
    posted_at=None,
    discovered_at=None,
) -> int:
    engine = get_engine(f"sqlite:///{db_path}")
    init_db(engine)
    factory = get_session_factory(engine)
    with factory() as session:
        source_id = None
        if source_name is not None:
            source = session.execute(
                select(JobSourceRow).where(JobSourceRow.name == source_name)
            ).scalar_one_or_none()
            if source is None:
                source = JobSourceRow(name=source_name, kind="ats_api", enabled=True)
                session.add(source)
                session.flush()
            source_id = source.id
        job = JobRow(
            source_id=source_id,
            company_name="Acme Corp",
            title="Business Analyst",
            location="Remote",
            remote_type="remote",
            employment_type="full_time",
            salary_min=80000,
            salary_max=110000,
            currency="USD",
            application_url="https://example.test/apply",
            job_fingerprint=fingerprint,
            description="Analyze business processes.",
            requirements="SQL, Excel",
            lifecycle_status=lifecycle_status,
            posted_at=posted_at,
        )
        session.add(job)
        session.flush()
        if discovered_at is not None:
            job.discovered_at = discovered_at
        job_id = job.id
        session.commit()
    return job_id


def _seed_match(
    db_path: Path,
    job_id: int,
    candidate_id: int,
    *,
    overall_score: int = 80,
    decision: str = "APPLY",
) -> None:
    engine = get_engine(f"sqlite:///{db_path}")
    factory = get_session_factory(engine)
    with factory() as session:
        session.add(
            JobMatch(
                job_id=job_id,
                candidate_id=candidate_id,
                overall_score=overall_score,
                decision=decision,
                skills_match=80,
                experience_match=70,
                role_match=90,
                project_match=60,
                education_match=100,
                location_match=100,
                seniority_match=80,
                eligibility_match=100,
                reasoning="Strong overlap with stated skills.",
                semantic_available=False,
            )
        )
        session.commit()


def _client(tmp_path: Path, monkeypatch, real_config) -> tuple[TestClient, Path]:
    db_path = _configure_env(tmp_path, monkeypatch, real_config)
    return TestClient(create_app()), db_path


def test_health(tmp_path, monkeypatch, real_config):
    client, _ = _client(tmp_path, monkeypatch, real_config)
    r = client.get("/api/health")
    assert r.status_code == 200
    assert r.json() == {"status": "ok"}


def test_candidate_profile_reflects_synthetic_files(tmp_path, monkeypatch, real_config):
    client, _ = _client(tmp_path, monkeypatch, real_config)
    r = client.get("/api/candidate/profile")
    assert r.status_code == 200
    body = r.json()
    assert body["name"] == SYNTHETIC_VALUES["full_name"]
    assert body["email"] == SYNTHETIC_VALUES["email"]
    assert "candidate_id" in body


def test_dashboard_summary_empty_by_default(tmp_path, monkeypatch, real_config):
    """No fabricated jobs/matches on a fresh install — everything at 0."""
    client, _ = _client(tmp_path, monkeypatch, real_config)
    r = client.get("/api/dashboard/summary")
    assert r.status_code == 200
    body = r.json()
    assert body["job_matches"] == 0
    assert body["total_jobs_discovered"] == 0
    assert body["top_opportunities"] == []


def test_jobs_scan_with_no_sources_enabled_is_honest_zero(tmp_path, monkeypatch, real_config):
    """config/sources.yaml ships with every source disabled — scanning
    must report zero results, never fabricate a job listing, and must not
    attempt any network call (build_sources returns an empty adapter
    list, so run_scan has nothing to iterate)."""
    client, _ = _client(tmp_path, monkeypatch, real_config)
    r = client.post("/api/jobs/scan")
    assert r.status_code == 200
    body = r.json()
    assert body["enabled_sources"] == 0
    assert body["results"] == []


def test_list_jobs_and_match_serialization(tmp_path, monkeypatch, real_config):
    client, db_path = _client(tmp_path, monkeypatch, real_config)
    profile_r = client.get("/api/candidate/profile")
    candidate_id = profile_r.json()["candidate_id"]

    job_id = _seed_job(db_path, fingerprint="job-1")
    _seed_match(db_path, job_id, candidate_id, overall_score=91, decision="APPLY")

    r = client.get("/api/jobs")
    assert r.status_code == 200
    body = r.json()
    assert body["total"] == 1
    item = body["items"][0]
    assert item["id"] == job_id
    assert item["match"]["overall_score"] == 91
    assert item["match"]["decision"] == "APPLY"
    assert item["pipeline_stage"] is None
    assert item["application_id"] is None


def test_job_out_carries_data_confidence_and_viability(tmp_path, monkeypatch, real_config):
    """Part 3.8/3.9: every JobOut carries confidence in the DATA (separate
    from the match score) and a practical application-viability read —
    computed honestly from the actual stored fields, not fabricated."""
    client, db_path = _client(tmp_path, monkeypatch, real_config)
    candidate_id = client.get("/api/candidate/profile").json()["candidate_id"]

    job_id = _seed_job(db_path, fingerprint="job-confidence")
    _seed_match(db_path, job_id, candidate_id, overall_score=91, decision="APPLY")

    item = client.get("/api/jobs").json()["items"][0]

    # _seed_job never sets posted_at or a company link, so confidence must
    # honestly report Medium — never a fabricated High next to those gaps.
    assert item["data_confidence"]["level"] == "Medium"
    assert any("posting date" in r.lower() for r in item["data_confidence"]["reasons"])
    assert any("company" in r.lower() for r in item["data_confidence"]["reasons"])

    # _seed_job has an application_url + ACTIVE lifecycle, and the match
    # above has no hard-stop/missing requirements and a perfect
    # location_match — so viability should be clean.
    assert item["viability"]["url_exists"] is True
    assert item["viability"]["job_active"] is True
    assert item["viability"]["qualifications_status"] == "MEETS"
    assert item["viability"]["location_compatible"] is True
    assert item["viability"]["overall"] == "VIABLE"


def test_check_url_endpoint_reports_live_reachability(tmp_path, monkeypatch, real_config):
    """The one live network check in viability — kept out of unit-test
    reach by monkeypatching the router's check_application_url so this
    test never makes a real request."""
    from datetime import datetime

    from job_agent.jobs.url_check import URLCheckResult
    from job_agent.web.routers import jobs as jobs_router

    client, db_path = _client(tmp_path, monkeypatch, real_config)
    job_id = _seed_job(db_path, fingerprint="job-url-check")

    monkeypatch.setattr(
        jobs_router,
        "check_application_url",
        lambda url, **_: URLCheckResult("REACHABLE", "HTTP 200", datetime.now(UTC)),
    )

    r = client.post(f"/api/jobs/{job_id}/check-url")
    assert r.status_code == 200
    body = r.json()
    assert body["status"] == "REACHABLE"
    assert body["detail"] == "HTTP 200"


def test_check_url_endpoint_404_for_missing_job(tmp_path, monkeypatch, real_config):
    client, _ = _client(tmp_path, monkeypatch, real_config)
    r = client.post("/api/jobs/999999/check-url")
    assert r.status_code == 404


def test_check_url_endpoint_400_when_no_application_url(tmp_path, monkeypatch, real_config):
    client, db_path = _client(tmp_path, monkeypatch, real_config)
    job_id = _seed_job(db_path, fingerprint="job-no-url")

    engine = get_engine(f"sqlite:///{db_path}")
    factory = get_session_factory(engine)
    with factory() as session:
        job = session.get(JobRow, job_id)
        job.application_url = None
        session.commit()

    r = client.post(f"/api/jobs/{job_id}/check-url")
    assert r.status_code == 400


def test_list_jobs_filters_by_min_score(tmp_path, monkeypatch, real_config):
    client, db_path = _client(tmp_path, monkeypatch, real_config)
    candidate_id = client.get("/api/candidate/profile").json()["candidate_id"]

    low_id = _seed_job(db_path, fingerprint="low")
    high_id = _seed_job(db_path, fingerprint="high")
    _seed_match(db_path, low_id, candidate_id, overall_score=20, decision="SKIP")
    _seed_match(db_path, high_id, candidate_id, overall_score=95, decision="APPLY")

    r = client.get("/api/jobs", params={"min_score": 50})
    assert r.status_code == 200
    ids = {item["id"] for item in r.json()["items"]}
    assert ids == {high_id}


def test_get_job_detail_404_for_missing_job(tmp_path, monkeypatch, real_config):
    client, _ = _client(tmp_path, monkeypatch, real_config)
    r = client.get("/api/jobs/999")
    assert r.status_code == 404


def test_compare_jobs_returns_full_detail_for_each(tmp_path, monkeypatch, real_config):
    """Part 3.10: comparing jobs must use the exact same JobDetailOut every
    other view uses — same match, confidence, and viability — never a
    second, separate comparison computation."""
    client, db_path = _client(tmp_path, monkeypatch, real_config)
    candidate_id = client.get("/api/candidate/profile").json()["candidate_id"]

    id_a = _seed_job(db_path, fingerprint="compare-a")
    id_b = _seed_job(db_path, fingerprint="compare-b")
    _seed_match(db_path, id_a, candidate_id, overall_score=91, decision="APPLY")
    _seed_match(db_path, id_b, candidate_id, overall_score=60, decision="REVIEW")

    r = client.get("/api/jobs/compare", params={"ids": f"{id_a},{id_b}"})
    assert r.status_code == 200
    body = r.json()
    assert {j["id"] for j in body} == {id_a, id_b}
    by_id = {j["id"]: j for j in body}
    assert by_id[id_a]["match"]["overall_score"] == 91
    assert by_id[id_b]["match"]["overall_score"] == 60
    assert "data_confidence" in by_id[id_a]
    assert "viability" in by_id[id_a]


def test_compare_jobs_requires_at_least_two_ids(tmp_path, monkeypatch, real_config):
    client, db_path = _client(tmp_path, monkeypatch, real_config)
    job_id = _seed_job(db_path, fingerprint="only-one")
    r = client.get("/api/jobs/compare", params={"ids": str(job_id)})
    assert r.status_code == 400


def test_compare_jobs_rejects_more_than_six(tmp_path, monkeypatch, real_config):
    client, db_path = _client(tmp_path, monkeypatch, real_config)
    ids = [_seed_job(db_path, fingerprint=f"compare-many-{i}") for i in range(7)]
    r = client.get("/api/jobs/compare", params={"ids": ",".join(str(i) for i in ids)})
    assert r.status_code == 400


def test_compare_jobs_404_when_any_id_missing(tmp_path, monkeypatch, real_config):
    client, db_path = _client(tmp_path, monkeypatch, real_config)
    job_id = _seed_job(db_path, fingerprint="compare-real")
    r = client.get("/api/jobs/compare", params={"ids": f"{job_id},999999"})
    assert r.status_code == 404


def test_compare_jobs_rejects_malformed_ids(tmp_path, monkeypatch, real_config):
    client, _ = _client(tmp_path, monkeypatch, real_config)
    r = client.get("/api/jobs/compare", params={"ids": "12,not-a-number"})
    assert r.status_code == 400


def test_save_job_creates_pipeline_entry_at_saved(tmp_path, monkeypatch, real_config):
    client, db_path = _client(tmp_path, monkeypatch, real_config)
    job_id = _seed_job(db_path, fingerprint="save-me")

    r = client.post(f"/api/jobs/{job_id}/save")
    assert r.status_code == 200
    body = r.json()
    assert body["pipeline_stage"] == "SAVED"
    assert body["application_id"] is not None

    pipeline_r = client.get("/api/pipeline")
    assert pipeline_r.status_code == 200
    items = pipeline_r.json()
    assert len(items) == 1
    assert items[0]["pipeline_stage"] == "SAVED"
    assert items[0]["job"]["id"] == job_id


def test_save_job_is_idempotent_and_never_regresses_stage(tmp_path, monkeypatch, real_config):
    client, db_path = _client(tmp_path, monkeypatch, real_config)
    job_id = _seed_job(db_path, fingerprint="idempotent")
    save_1 = client.post(f"/api/jobs/{job_id}/save").json()
    application_id = save_1["application_id"]

    client.patch(f"/api/pipeline/{application_id}", json={"pipeline_stage": "INTERVIEW"})

    save_2 = client.post(f"/api/jobs/{job_id}/save").json()
    assert save_2["application_id"] == application_id
    assert save_2["pipeline_stage"] == "INTERVIEW"  # never silently reset to SAVED


def test_pipeline_update_rejects_unknown_stage(tmp_path, monkeypatch, real_config):
    client, db_path = _client(tmp_path, monkeypatch, real_config)
    job_id = _seed_job(db_path, fingerprint="bad-stage")
    application_id = client.post(f"/api/jobs/{job_id}/save").json()["application_id"]

    r = client.patch(f"/api/pipeline/{application_id}", json={"pipeline_stage": "NOT_A_STAGE"})
    assert r.status_code == 422


def test_pipeline_update_notes_and_dates(tmp_path, monkeypatch, real_config):
    client, db_path = _client(tmp_path, monkeypatch, real_config)
    job_id = _seed_job(db_path, fingerprint="notes")
    application_id = client.post(f"/api/jobs/{job_id}/save").json()["application_id"]

    r = client.patch(
        f"/api/pipeline/{application_id}",
        json={
            "pipeline_stage": "INTERVIEW",
            "notes": "Recruiter call went well.",
            "recruiter_contact": "jane@acme.example",
            "interview_date": "2026-09-10T15:00:00Z",
        },
    )
    assert r.status_code == 200
    body = r.json()
    assert body["notes"] == "Recruiter call went well."
    assert body["recruiter_contact"] == "jane@acme.example"
    assert body["interview_date"].startswith("2026-09-10")


def test_pipeline_item_carries_scorecard_and_checklist(tmp_path, monkeypatch, real_config):
    """Part 4.11/4.12: the same scorecard/checklist logic used everywhere
    else — never a second, drifting computation on the pipeline view."""
    client, db_path = _client(tmp_path, monkeypatch, real_config)
    candidate_id = client.get("/api/candidate/profile").json()["candidate_id"]
    job_id = _seed_job(db_path, fingerprint="scorecard-job")
    _seed_match(db_path, job_id, candidate_id, overall_score=88, decision="APPLY")
    application_id = client.post(f"/api/jobs/{job_id}/save").json()["application_id"]

    body = client.get("/api/pipeline").json()[0]
    assert body["application_id"] == application_id
    assert body["scorecard"]["candidate_fit"] == 88.0
    assert body["scorecard"]["career_value"] == 100.0
    assert "overall_recommendation" in body["scorecard"]
    assert body["checklist"]["resume_selected"] is False
    assert body["checklist"]["cover_letter_ready"] is False
    assert body["checklist"]["submitted"] is False


def test_pipeline_checklist_toggles_are_manual_and_persist(tmp_path, monkeypatch, real_config):
    client, db_path = _client(tmp_path, monkeypatch, real_config)
    job_id = _seed_job(db_path, fingerprint="checklist-job")
    application_id = client.post(f"/api/jobs/{job_id}/save").json()["application_id"]

    r = client.patch(
        f"/api/pipeline/{application_id}",
        json={"cover_letter_ready": True, "questions_prepared": True},
    )
    assert r.status_code == 200
    assert r.json()["checklist"]["cover_letter_ready"] is True
    assert r.json()["checklist"]["questions_prepared"] is True

    # Persisted, not just echoed back.
    refetched = client.get("/api/pipeline").json()[0]
    assert refetched["checklist"]["cover_letter_ready"] is True
    assert refetched["checklist"]["questions_prepared"] is True


def test_pipeline_history_records_every_change_never_overwrites(tmp_path, monkeypatch, real_config):
    """Part 4.13: History is append-only — two edits produce two entries,
    and the first entry's content is never lost."""
    client, db_path = _client(tmp_path, monkeypatch, real_config)
    job_id = _seed_job(db_path, fingerprint="history-job")
    application_id = client.post(f"/api/jobs/{job_id}/save").json()["application_id"]

    client.patch(f"/api/pipeline/{application_id}", json={"notes": "First note."})
    client.patch(f"/api/pipeline/{application_id}", json={"notes": "Second note."})

    r = client.get(f"/api/pipeline/{application_id}/history")
    assert r.status_code == 200
    events = r.json()
    # At least the initial APPLICATION_DISCOVERED bookkeeping event plus
    # two PIPELINE_UPDATED entries from the two distinct note edits above.
    pipeline_updates = [e for e in events if e["event_type"] == "PIPELINE_UPDATED"]
    assert len(pipeline_updates) == 2
    assert pipeline_updates[0]["details"]["notes"]["to"] == "First note."
    assert pipeline_updates[1]["details"]["notes"]["to"] == "Second note."
    # Oldest first.
    assert events[0]["created_at"] <= events[-1]["created_at"]


def test_pipeline_history_404_for_missing_application(tmp_path, monkeypatch, real_config):
    client, _ = _client(tmp_path, monkeypatch, real_config)
    r = client.get("/api/pipeline/999999/history")
    assert r.status_code == 404


def test_pipeline_update_with_no_changes_appends_no_event(tmp_path, monkeypatch, real_config):
    client, db_path = _client(tmp_path, monkeypatch, real_config)
    job_id = _seed_job(db_path, fingerprint="no-op-job")
    application_id = client.post(f"/api/jobs/{job_id}/save").json()["application_id"]

    before = len(client.get(f"/api/pipeline/{application_id}/history").json())
    # Re-sending the SAME pipeline_stage it already has is not a change.
    client.patch(f"/api/pipeline/{application_id}", json={"pipeline_stage": "SAVED"})
    after = len(client.get(f"/api/pipeline/{application_id}/history").json())
    assert after == before


def test_pipeline_delete_refuses_past_shortlisted(tmp_path, monkeypatch, real_config):
    client, db_path = _client(tmp_path, monkeypatch, real_config)
    job_id = _seed_job(db_path, fingerprint="protect-history")
    application_id = client.post(f"/api/jobs/{job_id}/save").json()["application_id"]
    client.patch(f"/api/pipeline/{application_id}", json={"pipeline_stage": "APPLIED"})

    r = client.delete(f"/api/pipeline/{application_id}")
    assert r.status_code == 409

    still_there = client.get("/api/pipeline").json()
    assert len(still_there) == 1


def test_pipeline_delete_allowed_while_saved(tmp_path, monkeypatch, real_config):
    client, db_path = _client(tmp_path, monkeypatch, real_config)
    job_id = _seed_job(db_path, fingerprint="ok-to-remove")
    application_id = client.post(f"/api/jobs/{job_id}/save").json()["application_id"]

    r = client.delete(f"/api/pipeline/{application_id}")
    assert r.status_code == 204


def test_follow_ups_empty_for_fresh_application(tmp_path, monkeypatch, real_config):
    client, db_path = _client(tmp_path, monkeypatch, real_config)
    job_id = _seed_job(db_path, fingerprint="fresh-app")
    application_id = client.post(f"/api/jobs/{job_id}/save").json()["application_id"]
    client.patch(f"/api/pipeline/{application_id}", json={"pipeline_stage": "APPLIED"})

    r = client.get("/api/pipeline/follow-ups")
    assert r.status_code == 200
    assert r.json() == []


def test_follow_ups_surfaces_stale_applied_application(tmp_path, monkeypatch, real_config):
    from datetime import UTC, datetime, timedelta

    from job_agent.db.models import Application

    client, db_path = _client(tmp_path, monkeypatch, real_config)
    job_id = _seed_job(db_path, fingerprint="stale-app")
    application_id = client.post(f"/api/jobs/{job_id}/save").json()["application_id"]
    client.patch(f"/api/pipeline/{application_id}", json={"pipeline_stage": "APPLIED"})

    engine = get_engine(f"sqlite:///{db_path}")
    with get_session_factory(engine)() as session:
        application = session.get(Application, application_id)
        application.updated_at = datetime.now(UTC) - timedelta(days=10)
        session.commit()

    r = client.get("/api/pipeline/follow-ups")
    assert r.status_code == 200
    body = r.json()
    assert len(body) == 1
    assert body[0]["application_id"] == application_id
    assert body[0]["applied_days_ago"] >= 10
    assert "follow-up" in body[0]["suggested_action"].lower()


def test_analytics_reflects_pipeline_stage_counts(tmp_path, monkeypatch, real_config):
    client, db_path = _client(tmp_path, monkeypatch, real_config)
    job_a = _seed_job(db_path, fingerprint="analytics-a")
    job_b = _seed_job(db_path, fingerprint="analytics-b")
    app_a = client.post(f"/api/jobs/{job_a}/save").json()["application_id"]
    app_b = client.post(f"/api/jobs/{job_b}/save").json()["application_id"]
    client.patch(f"/api/pipeline/{app_a}", json={"pipeline_stage": "INTERVIEW"})
    client.patch(f"/api/pipeline/{app_b}", json={"pipeline_stage": "APPLIED"})

    r = client.get("/api/pipeline/analytics")
    assert r.status_code == 200
    body = r.json()
    stage_counts = {row["stage"]: row["count"] for row in body["stage_breakdown"]}
    assert stage_counts["INTERVIEW"] == 1
    assert stage_counts["APPLIED"] == 1
    assert body["total_interviews"] == 1


def test_jobs_match_runs_deterministic_only_without_llm_key(tmp_path, monkeypatch, real_config):
    """No ANTHROPIC_API_KEY configured -> NullLLMProvider -> deterministic
    matching only, exactly like `jobs match` from the CLI; must not raise."""
    client, db_path = _client(tmp_path, monkeypatch, real_config)
    _seed_job(db_path, fingerprint="match-me")

    r = client.post("/api/jobs/match")
    assert r.status_code == 200
    body = r.json()
    assert body["matched"] == 1


def test_resume_upload_rejects_non_docx(tmp_path, monkeypatch, real_config):
    client, _ = _client(tmp_path, monkeypatch, real_config)
    r = client.post(
        "/api/candidate/resume",
        files={"file": ("resume.txt", b"not a docx", "text/plain")},
    )
    assert r.status_code == 422


def test_resume_upload_accepts_valid_docx(tmp_path, monkeypatch, real_config):
    import io

    import docx

    client, _ = _client(tmp_path, monkeypatch, real_config)
    buf = io.BytesIO()
    document = docx.Document()
    document.add_paragraph("Uploaded synthetic resume content for testing.")
    document.save(buf)
    buf.seek(0)

    r = client.post(
        "/api/candidate/resume",
        files={
            "file": (
                "resume.docx",
                buf,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert r.status_code == 200
    body = r.json()
    assert body["saved_path"].endswith("resume_master.docx")
    assert body["validation_status"] in ("PASSED", "FAILED")


def test_spa_fallback_serves_index_for_client_routes(tmp_path, monkeypatch, real_config):
    """React Router uses client-side (BrowserRouter) routing — a direct
    GET for a client-only path like /pipeline or /jobs/5 (a bookmark, a
    refresh) must still return the SPA shell, not 404, whenever the
    frontend has been built (`web-ui/dist` exists in this repo checkout).
    An unknown /api/* path must still 404 as a real API 404, never be
    swallowed by the SPA fallback."""
    from job_agent.web.app import _FRONTEND_DIST

    if not _FRONTEND_DIST.exists():
        import pytest

        pytest.skip("web-ui/dist not built in this checkout — run `npm run build` in web-ui/")

    client, _ = _client(tmp_path, monkeypatch, real_config)
    for path in ("/pipeline", "/jobs/5", "/resume", "/analytics"):
        r = client.get(path)
        assert r.status_code == 200, path
        assert '<div id="root">' in r.text

    r = client.get("/api/this-route-does-not-exist")
    assert r.status_code == 404


def test_duplicate_jobs_across_sources_show_as_one_canonical_entry(
    tmp_path, monkeypatch, real_config
):
    """Two Job rows with the SAME fingerprint (same content, different
    sources) must appear as ONE card, with the other source listed under
    also_seen_on — never two separate cards (section 8)."""
    from datetime import UTC, datetime, timedelta

    client, db_path = _client(tmp_path, monkeypatch, real_config)
    older = datetime.now(UTC) - timedelta(days=2)
    newer = datetime.now(UTC) - timedelta(hours=1)
    _seed_job(
        db_path,
        fingerprint="dup-fp",
        source_name="greenhouse",
        posted_at=older,
    )
    _seed_job(
        db_path,
        fingerprint="dup-fp",
        source_name="lever",
        posted_at=newer,
    )

    r = client.get("/api/jobs")
    assert r.status_code == 200
    body = r.json()
    assert body["total"] == 1
    item = body["items"][0]
    assert item["source_name"] == "lever"  # the freshest of the two
    assert item["also_seen_on"] == ["greenhouse"]
    assert item["duplicate_count"] == 1


def test_inactive_jobs_excluded_by_default_included_on_request(tmp_path, monkeypatch, real_config):
    client, db_path = _client(tmp_path, monkeypatch, real_config)
    _seed_job(db_path, fingerprint="active-job", lifecycle_status="ACTIVE")
    _seed_job(db_path, fingerprint="closed-job", lifecycle_status="CLOSED")
    _seed_job(db_path, fingerprint="expired-job", lifecycle_status="EXPIRED")

    default_r = client.get("/api/jobs")
    assert default_r.json()["total"] == 1

    all_r = client.get("/api/jobs", params={"include_inactive": "true"})
    assert all_r.json()["total"] == 3


def test_top10_ranks_by_composite_score_not_match_alone(tmp_path, monkeypatch, real_config):
    client, db_path = _client(tmp_path, monkeypatch, real_config)
    candidate_id = client.get("/api/candidate/profile").json()["candidate_id"]

    weak_job = _seed_job(db_path, fingerprint="weak")
    strong_job = _seed_job(db_path, fingerprint="strong")
    _seed_match(db_path, weak_job, candidate_id, overall_score=30, decision="SKIP")
    _seed_match(db_path, strong_job, candidate_id, overall_score=95, decision="APPLY")

    r = client.get("/api/jobs/top10")
    assert r.status_code == 200
    body = r.json()
    assert len(body) == 2
    assert body[0]["job"]["id"] == strong_job
    assert body[0]["rank_score"] >= body[1]["rank_score"]
    assert body[0]["recommendation"] == "Apply today."


def test_top10_excludes_closed_and_expired_jobs(tmp_path, monkeypatch, real_config):
    client, db_path = _client(tmp_path, monkeypatch, real_config)
    candidate_id = client.get("/api/candidate/profile").json()["candidate_id"]

    active_job = _seed_job(db_path, fingerprint="active")
    closed_job = _seed_job(db_path, fingerprint="closed", lifecycle_status="CLOSED")
    _seed_match(db_path, active_job, candidate_id, overall_score=90, decision="APPLY")
    _seed_match(db_path, closed_job, candidate_id, overall_score=99, decision="APPLY")

    r = client.get("/api/jobs/top10")
    ids = [item["job"]["id"] for item in r.json()]
    assert active_job in ids
    assert closed_job not in ids


def test_apply_now_queue_only_includes_apply_decisions(tmp_path, monkeypatch, real_config):
    client, db_path = _client(tmp_path, monkeypatch, real_config)
    candidate_id = client.get("/api/candidate/profile").json()["candidate_id"]

    apply_job = _seed_job(db_path, fingerprint="apply-me")
    review_job = _seed_job(db_path, fingerprint="review-me")
    _seed_match(db_path, apply_job, candidate_id, overall_score=95, decision="APPLY")
    _seed_match(db_path, review_job, candidate_id, overall_score=70, decision="REVIEW")

    r = client.get("/api/jobs/apply-now")
    assert r.status_code == 200
    ids = [item["job"]["id"] for item in r.json()]
    assert apply_job in ids
    assert review_job not in ids


def test_tailor_resume_returns_deterministic_result_with_no_api_key(
    tmp_path, monkeypatch, real_config
):
    client, db_path = _client(tmp_path, monkeypatch, real_config)
    job_id = _seed_job(db_path, fingerprint="tailor-me")

    r = client.post(f"/api/jobs/{job_id}/tailor-resume")
    assert r.status_code == 200
    body = r.json()
    assert body["job_id"] == job_id
    assert body["generated_by"] == "deterministic"
    assert body["professional_summary"]
    assert isinstance(body["relevant_skills"], list)
    assert isinstance(body["ats_keywords"], list)


def test_tailor_resume_404_for_unknown_job(tmp_path, monkeypatch, real_config):
    client, _ = _client(tmp_path, monkeypatch, real_config)
    r = client.post("/api/jobs/999999/tailor-resume")
    assert r.status_code == 404


def test_cover_letter_is_job_specific_and_deterministic_with_no_api_key(
    tmp_path, monkeypatch, real_config
):
    client, db_path = _client(tmp_path, monkeypatch, real_config)
    job_id = _seed_job(db_path, fingerprint="cover-letter-me")

    r = client.post(f"/api/jobs/{job_id}/cover-letter")
    assert r.status_code == 200
    body = r.json()
    assert body["job_id"] == job_id
    assert body["generated_by"] == "deterministic"
    assert "Business Analyst" in body["body"]
    assert "Acme Corp" in body["body"]


def test_cover_letter_404_for_unknown_job(tmp_path, monkeypatch, real_config):
    client, _ = _client(tmp_path, monkeypatch, real_config)
    r = client.post("/api/jobs/999999/cover-letter")
    assert r.status_code == 404


def test_assistant_grounds_answers_in_candidate_profile(tmp_path, monkeypatch, real_config):
    client, db_path = _client(tmp_path, monkeypatch, real_config)
    job_id = _seed_job(db_path, fingerprint="assistant-me")

    r = client.post(
        f"/api/jobs/{job_id}/assistant",
        json={"questions": ["What is your email address?", "What is your expected salary?"]},
    )
    assert r.status_code == 200
    body = r.json()
    assert body["job_id"] == job_id
    assert len(body["answers"]) == 2
    email_answer = body["answers"][0]
    assert email_answer["answer"] == SYNTHETIC_VALUES["email"]
    assert email_answer["requires_human"] is False
    salary_answer = body["answers"][1]
    assert salary_answer["answer"] is None
    assert salary_answer["requires_human"] is True


def test_assistant_404_for_unknown_job(tmp_path, monkeypatch, real_config):
    client, _ = _client(tmp_path, monkeypatch, real_config)
    r = client.post("/api/jobs/999999/assistant", json={"questions": ["What is your email?"]})
    assert r.status_code == 404


def test_assistant_rejects_empty_question_list(tmp_path, monkeypatch, real_config):
    client, db_path = _client(tmp_path, monkeypatch, real_config)
    job_id = _seed_job(db_path, fingerprint="assistant-empty")
    r = client.post(f"/api/jobs/{job_id}/assistant", json={"questions": []})
    assert r.status_code == 422


def test_companies_empty_by_default(tmp_path, monkeypatch, real_config):
    client, _ = _client(tmp_path, monkeypatch, real_config)
    r = client.get("/api/companies")
    assert r.status_code == 200
    assert r.json() == []


def test_companies_aggregates_jobs_by_company_name(tmp_path, monkeypatch, real_config):
    client, db_path = _client(tmp_path, monkeypatch, real_config)
    job_a = _seed_job(db_path, fingerprint="company-a-1")
    job_b = _seed_job(db_path, fingerprint="company-a-2")

    r = client.get("/api/companies")
    assert r.status_code == 200
    body = r.json()
    assert len(body) == 1
    assert body[0]["name"] == "Acme Corp"
    assert body[0]["open_roles"] == 2
    job_ids = {j["id"] for j in body[0]["matching_jobs"]}
    assert job_ids == {job_a, job_b}
    # No verified company data exists — never fabricated.
    assert body[0]["industry"] is None
    assert body[0]["size"] is None


def test_company_fit_unknown_with_no_matches(tmp_path, monkeypatch, real_config):
    client, db_path = _client(tmp_path, monkeypatch, real_config)
    _seed_job(db_path, fingerprint="no-match-co")

    r = client.get("/api/companies")
    body = r.json()
    assert body[0]["company_fit"] is None


def test_company_fit_reflects_real_match_scores(tmp_path, monkeypatch, real_config):
    client, db_path = _client(tmp_path, monkeypatch, real_config)
    candidate_id = client.get("/api/candidate/profile").json()["candidate_id"]
    job_id = _seed_job(db_path, fingerprint="high-fit-co")
    _seed_match(db_path, job_id, candidate_id, overall_score=95, decision="APPLY")

    r = client.get("/api/companies")
    body = r.json()
    assert body[0]["company_fit"] is not None
    assert body[0]["company_fit"] >= 80


def test_get_company_by_id_returns_company(tmp_path, monkeypatch, real_config):
    client, db_path = _client(tmp_path, monkeypatch, real_config)
    job_id = _seed_job(db_path, fingerprint="single-co")

    r = client.get(f"/api/companies/{job_id}")
    assert r.status_code == 200
    assert r.json()["name"] == "Acme Corp"


def test_get_company_404_for_unknown_id(tmp_path, monkeypatch, real_config):
    client, _ = _client(tmp_path, monkeypatch, real_config)
    r = client.get("/api/companies/999999")
    assert r.status_code == 404


def test_watchlist_empty_by_default(tmp_path, monkeypatch, real_config):
    client, _ = _client(tmp_path, monkeypatch, real_config)
    r = client.get("/api/watchlist")
    assert r.status_code == 200
    assert r.json() == []


def test_watchlist_add_list_delete(tmp_path, monkeypatch, real_config):
    client, _ = _client(tmp_path, monkeypatch, real_config)

    r = client.post("/api/watchlist", json={"kind": "COMPANY", "value": "Acme Corp"})
    assert r.status_code == 200
    entry_id = r.json()["id"]
    assert r.json()["kind"] == "COMPANY"
    assert r.json()["value"] == "Acme Corp"

    r = client.get("/api/watchlist")
    assert len(r.json()) == 1

    r = client.delete(f"/api/watchlist/{entry_id}")
    assert r.status_code == 204

    r = client.get("/api/watchlist")
    assert r.json() == []


def test_watchlist_rejects_invalid_kind(tmp_path, monkeypatch, real_config):
    client, _ = _client(tmp_path, monkeypatch, real_config)
    r = client.post("/api/watchlist", json={"kind": "INVALID", "value": "x"})
    assert r.status_code == 422


def test_watchlist_add_is_idempotent_for_same_kind_and_value(tmp_path, monkeypatch, real_config):
    client, _ = _client(tmp_path, monkeypatch, real_config)
    client.post("/api/watchlist", json={"kind": "ROLE", "value": "Business Analyst"})
    client.post("/api/watchlist", json={"kind": "ROLE", "value": "Business Analyst"})
    r = client.get("/api/watchlist")
    assert len(r.json()) == 1


def test_watchlist_delete_404_for_unknown_entry(tmp_path, monkeypatch, real_config):
    client, _ = _client(tmp_path, monkeypatch, real_config)
    r = client.delete("/api/watchlist/999999")
    assert r.status_code == 404


def test_notifications_empty_by_default(tmp_path, monkeypatch, real_config):
    client, _ = _client(tmp_path, monkeypatch, real_config)
    r = client.get("/api/notifications")
    assert r.status_code == 200
    assert r.json() == []


def test_notifications_generated_from_high_match_and_marked_read(
    tmp_path, monkeypatch, real_config
):
    client, db_path = _client(tmp_path, monkeypatch, real_config)
    candidate_id = client.get("/api/candidate/profile").json()["candidate_id"]
    job_id = _seed_job(db_path, fingerprint="notif-me")
    _seed_match(db_path, job_id, candidate_id, overall_score=95, decision="APPLY")

    engine = get_engine(f"sqlite:///{db_path}")
    with get_session_factory(engine)() as session:
        job = session.get(JobRow, job_id)
        match = session.execute(select(JobMatch).where(JobMatch.job_id == job_id)).scalar_one()
        generate_notifications(session, candidate_id, [(job, match)], min_score=90)

    r = client.get("/api/notifications")
    assert r.status_code == 200
    body = r.json()
    assert len(body) == 1
    assert body[0]["event_type"] == "HIGH_MATCH"
    assert body[0]["read_at"] is None

    notif_id = body[0]["id"]
    r = client.post(f"/api/notifications/{notif_id}/read")
    assert r.status_code == 200
    assert r.json()["read_at"] is not None

    r = client.get("/api/notifications", params={"unread_only": "true"})
    assert r.json() == []


def test_notifications_mark_read_404_for_unknown(tmp_path, monkeypatch, real_config):
    client, _ = _client(tmp_path, monkeypatch, real_config)
    r = client.post("/api/notifications/999999/read")
    assert r.status_code == 404


def test_search_preferences_default_when_no_row(tmp_path, monkeypatch, real_config):
    client, _ = _client(tmp_path, monkeypatch, real_config)
    r = client.get("/api/settings/search-preferences")
    assert r.status_code == 200
    body = r.json()
    assert body["target_roles"] == []
    assert body["min_match_score"] == 0
    assert body["search_frequency_hours"] == 24
    assert body["notification_min_score"] == 90
    assert body["notification_frequency"] == "daily"


def test_search_preferences_update_persists_and_merges(tmp_path, monkeypatch, real_config):
    client, _ = _client(tmp_path, monkeypatch, real_config)

    r = client.put(
        "/api/settings/search-preferences",
        json={"target_roles": ["Business Analyst"], "min_match_score": 70},
    )
    assert r.status_code == 200
    assert r.json()["target_roles"] == ["Business Analyst"]
    assert r.json()["min_match_score"] == 70

    # A second, partial update must not reset the first update's fields.
    r = client.put("/api/settings/search-preferences", json={"notification_frequency": "weekly"})
    assert r.status_code == 200
    body = r.json()
    assert body["target_roles"] == ["Business Analyst"]
    assert body["min_match_score"] == 70
    assert body["notification_frequency"] == "weekly"

    r = client.get("/api/settings/search-preferences")
    assert r.json()["target_roles"] == ["Business Analyst"]


def test_supported_countries_lists_expected_set(tmp_path, monkeypatch, real_config):
    client, _ = _client(tmp_path, monkeypatch, real_config)
    r = client.get("/api/settings/supported-countries")
    assert r.status_code == 200
    assert "India" in r.json()
    assert "Remote" in r.json()


def test_insights_empty_with_no_matched_jobs(tmp_path, monkeypatch, real_config):
    client, _ = _client(tmp_path, monkeypatch, real_config)
    r = client.get("/api/candidate/insights")
    assert r.status_code == 200
    assert r.json() == {"categories": [], "summary": []}


def test_insights_surfaces_notable_pattern_from_saved_jobs(tmp_path, monkeypatch, real_config):
    client, db_path = _client(tmp_path, monkeypatch, real_config)
    candidate_id = client.get("/api/candidate/profile").json()["candidate_id"]

    saved_ids = []
    for i in range(4):
        job_id = _seed_job(db_path, fingerprint=f"remote-saved-{i}")
        _seed_match(db_path, job_id, candidate_id, overall_score=80)
        client.post(f"/api/jobs/{job_id}/save")
        saved_ids.append(job_id)

    engine = get_engine(f"sqlite:///{db_path}")
    with get_session_factory(engine)() as session:
        for job_id in saved_ids:
            job = session.get(JobRow, job_id)
            job.remote_type = "remote"
        session.commit()

    for i in range(4):
        job_id = _seed_job(db_path, fingerprint=f"onsite-ignored-{i}")
        _seed_match(db_path, job_id, candidate_id, overall_score=80)
        with get_session_factory(engine)() as session:
            job = session.get(JobRow, job_id)
            job.remote_type = "onsite"
            session.commit()

    r = client.get("/api/candidate/insights")
    assert r.status_code == 200
    body = r.json()
    assert any("remote" in s.lower() for s in body["summary"])
    remote_category = next(c for c in body["categories"] if c["category"] == "remote")
    assert remote_category["saved"] == 4


def test_new_since_last_visit_shows_everything_on_first_visit(tmp_path, monkeypatch, real_config):
    client, db_path = _client(tmp_path, monkeypatch, real_config)
    _seed_job(db_path, fingerprint="first-visit-job")

    r = client.get("/api/dashboard/new-since-last-visit")
    assert r.status_code == 200
    body = r.json()
    assert body["previous_visit_at"] is None
    assert len(body["jobs"]) == 1


def test_new_since_last_visit_never_repeats_the_same_job(tmp_path, monkeypatch, real_config):
    client, db_path = _client(tmp_path, monkeypatch, real_config)
    _seed_job(db_path, fingerprint="repeat-check-job")

    first = client.get("/api/dashboard/new-since-last-visit")
    assert len(first.json()["jobs"]) == 1

    second = client.get("/api/dashboard/new-since-last-visit")
    assert second.json()["previous_visit_at"] is not None
    assert second.json()["jobs"] == []


def test_new_since_last_visit_shows_only_jobs_discovered_after_previous_visit(
    tmp_path, monkeypatch, real_config
):
    from datetime import UTC, datetime, timedelta

    client, db_path = _client(tmp_path, monkeypatch, real_config)
    _seed_job(db_path, fingerprint="old-job")
    client.get("/api/dashboard/new-since-last-visit")

    fresh_job_id = _seed_job(
        db_path, fingerprint="fresh-job", discovered_at=datetime.now(UTC) + timedelta(minutes=1)
    )

    r = client.get("/api/dashboard/new-since-last-visit")
    body = r.json()
    assert len(body["jobs"]) == 1
    assert body["jobs"][0]["job"]["id"] == fresh_job_id


def test_career_profile_grounded_in_real_career_paths(tmp_path, monkeypatch, real_config):
    client, _ = _client(tmp_path, monkeypatch, real_config)
    r = client.get("/api/candidate/career-profile")
    assert r.status_code == 200
    body = r.json()
    assert "primary_direction" in body
    assert "best_locations" in body
    # Never fabricated when the synthetic candidate has no real skills/paths.
    if body["primary_direction"] is None:
        assert body["strengths"] == []


def test_career_path_comparison_empty_when_no_paths(tmp_path, monkeypatch, real_config):
    client, _ = _client(tmp_path, monkeypatch, real_config)
    r = client.get("/api/candidate/career-paths/compare")
    assert r.status_code == 200
    assert isinstance(r.json(), list)


def test_skill_gaps_empty_by_default(tmp_path, monkeypatch, real_config):
    client, _ = _client(tmp_path, monkeypatch, real_config)
    r = client.get("/api/candidate/skill-gaps")
    assert r.status_code == 200
    assert r.json() == []


def test_morning_briefing_empty_by_default(tmp_path, monkeypatch, real_config):
    client, _ = _client(tmp_path, monkeypatch, real_config)
    r = client.get("/api/dashboard/briefing")
    assert r.status_code == 200
    body = r.json()
    assert body["total_opportunities"] == 0
    assert body["top_highlights"] == []
    assert body["follow_up_summaries"] == []


def test_morning_briefing_counts_and_highlights_reflect_real_matches(
    tmp_path, monkeypatch, real_config
):
    client, db_path = _client(tmp_path, monkeypatch, real_config)
    candidate_id = client.get("/api/candidate/profile").json()["candidate_id"]
    apply_job = _seed_job(db_path, fingerprint="briefing-apply")
    review_job = _seed_job(db_path, fingerprint="briefing-review")
    _seed_match(db_path, apply_job, candidate_id, overall_score=95, decision="APPLY")
    _seed_match(db_path, review_job, candidate_id, overall_score=70, decision="REVIEW")

    r = client.get("/api/dashboard/briefing")
    assert r.status_code == 200
    body = r.json()
    assert body["exceptional_count"] == 1
    assert body["strong_count"] == 1
    assert body["total_opportunities"] == 2
    assert len(body["top_highlights"]) >= 1
    assert body["top_highlights"][0]["job_id"] == apply_job
