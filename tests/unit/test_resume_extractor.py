from __future__ import annotations

import pytest

from job_agent.resume.errors import ResumeExtractionError
from job_agent.resume.extractor import extract_resume_text


def test_extracts_real_resume(repo_root):
    text = extract_resume_text(repo_root / "candidate" / "resume_master.docx")
    assert "Priyanshu Vats".lower() in text.lower() or "PRIYANSHU VATS" in text
    assert "Instawork Robotics Labs" in text
    assert len(text) > 500


def test_missing_file_raises(tmp_path):
    with pytest.raises(ResumeExtractionError, match="not found"):
        extract_resume_text(tmp_path / "does_not_exist.docx")


def test_non_docx_file_raises(tmp_path):
    bad = tmp_path / "resume.docx"
    bad.write_text("this is not a real docx file, just plain text")
    with pytest.raises(ResumeExtractionError, match="Could not open"):
        extract_resume_text(bad)


def test_empty_docx_raises(tmp_path):
    import docx

    empty = tmp_path / "empty.docx"
    document = docx.Document()
    document.save(str(empty))
    with pytest.raises(ResumeExtractionError, match="no extractable text"):
        extract_resume_text(empty)


def test_extracts_table_text(tmp_path):
    import docx

    path = tmp_path / "with_table.docx"
    document = docx.Document()
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Skill"
    table.rows[0].cells[1].text = "Python"
    document.save(str(path))

    text = extract_resume_text(path)
    assert "Python" in text


def test_missing_python_docx_raises(tmp_path, monkeypatch):
    import builtins

    real_import = builtins.__import__

    def fake_import(name, *args, **kwargs):
        if name == "docx":
            raise ImportError("simulated missing dependency")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", fake_import)
    existing = tmp_path / "resume.docx"
    existing.write_bytes(b"placeholder")
    with pytest.raises(ResumeExtractionError, match="python-docx is not installed"):
        extract_resume_text(existing)
