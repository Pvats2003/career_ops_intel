"""First real-execution-vertical-slice checkpoint — CLI integration tests
for `applications browser-approve` and `applications browser-fill`: the
approval-gated fill boundary between `applications browser-preview`
(pure preview, creates nothing) and a future submission checkpoint
(NOT implemented here — `submit()` remains structurally refused).

Same conventions as `test_cli_browser_preview.py`: `typer.testing.
CliRunner`, in-process, temporary sqlite DB + config copy, every
live-network branch points at a LOCAL synthetic fixture on 127.0.0.1,
`CANDIDATE_DIR` always points at a minimal synthetic candidate directory
(never the real one), and every candidate value that could appear in
`result.output` is one of the SYNTHETIC_VALUES placeholders below.

State-transition coverage (Step 10):
    UNPREPARED -> PREPARED -> HUMAN_REVIEW -> APPROVED -> FILL -> FILLED -> STOP
maps onto this codebase's EXISTING primitives rather than new ones:
    UNPREPARED      = no approval / no prior browser-fill for this application
    PREPARED        = _run_browser_preparation() has produced a HumanReviewSnapshot
    HUMAN_REVIEW    = the snapshot rendered to the human (browser-approve's output)
    APPROVED        = a valid ApplicationApproval row exists, bound to the exact
                       job fingerprint + snapshot fingerprint (Phase 6C machinery,
                       completely unmodified)
    FILL            = browser-fill re-runs preparation in a FRESH browser session
    FILLED          = an ApplicationEvent "BROWSER_FILL_COMPLETED" is recorded
    STOP            = the command returns; submit() is never called anywhere
"""

from __future__ import annotations

import shutil
from pathlib import Path

import playwright.sync_api as playwright_sync_api
import pytest
from sqlalchemy import select
from typer.testing import CliRunner

from job_agent.candidate.schema import (
    CandidateProfile,
    Fact,
    LocationPreferences,
    SalaryPreferences,
    TargetRoles,
    VisaInformation,
    WorkPreferences,
)
from job_agent.cli.main import app
from job_agent.db.models import (
    Application,
    ApplicationApproval,
    ApplicationEvent,
    Company,
    JobSource,
)
from job_agent.db.models import Job as JobRow
from job_agent.db.repository import save_candidate_profile
from job_agent.db.session import get_engine, get_session_factory, init_db

pytest.importorskip("playwright.sync_api")

CHROMIUM_EXECUTABLE = "/opt/pw-browsers/chromium-1194/chrome-linux/chrome"

pytestmark = pytest.mark.skipif(
    not Path(CHROMIUM_EXECUTABLE).exists(),
    reason="pre-installed Chromium not available in this environment",
)

runner = CliRunner()

SYNTHETIC_VALUES = {
    "full_name": "Test Candidate",
    "email": "test.candidate@example.invalid",
    "phone": "+1-555-0100",
    "current_location": "Test City, Test Country",
    "linkedin_url": "https://linkedin.com/in/testcandidate",
    "current_company": "Synthetic Employer Inc.",
}


def _fact(value: str) -> Fact[str]:
    return Fact[str](
        value=value, source="test_synthetic_profile", confidence=1.0, verified=True
    )


def _synthetic_identity_profile() -> CandidateProfile:
    unknown_pref = Fact.unknown(source="test_synthetic_profile")
    return CandidateProfile(
        identity_name=_fact(SYNTHETIC_VALUES["full_name"]),
        identity_current_location=_fact(SYNTHETIC_VALUES["current_location"]),
        contact_email=_fact(SYNTHETIC_VALUES["email"]),
        contact_phone=_fact(SYNTHETIC_VALUES["phone"]),
        contact_linkedin=_fact(SYNTHETIC_VALUES["linkedin_url"]),
        target_roles=TargetRoles(),
        work_preferences=WorkPreferences(
            remote=unknown_pref, willing_to_relocate=unknown_pref, notice_period=unknown_pref
        ),
        location_preferences=LocationPreferences(
            current_location=_fact(SYNTHETIC_VALUES["current_location"]),
            open_to_countries=unknown_pref,
        ),
        salary_preferences=SalaryPreferences(
            currency=unknown_pref, minimum_annual=unknown_pref,
            target_annual=unknown_pref, negotiable=unknown_pref,
        ),
        visa_information=VisaInformation(
            nationality=unknown_pref, requires_sponsorship_us=unknown_pref,
            requires_sponsorship_uk=unknown_pref, requires_sponsorship_eu=unknown_pref,
            requires_sponsorship_other=unknown_pref,
        ),
    )


def _make_synthetic_candidate_dir(candidate_dir: Path) -> Path:
    candidate_dir.mkdir(parents=True, exist_ok=True)
    (candidate_dir / "profile.md").write_text(
        "## Identity\n"
        f"name: {SYNTHETIC_VALUES['full_name']}\n"
        f"current_location: {SYNTHETIC_VALUES['current_location']}\n"
        "\n"
        "## Contact\n"
        f"email: {SYNTHETIC_VALUES['email']}\n"
        f"phone: {SYNTHETIC_VALUES['phone']}\n"
        f"linkedin: {SYNTHETIC_VALUES['linkedin_url']}\n"
    )
    for filename in (
        "experience.md", "projects.md", "skills.md", "education.md", "achievements.md",
    ):
        (candidate_dir / filename).write_text("")
    (candidate_dir / "answers").mkdir(exist_ok=True)

    import docx

    document = docx.Document()
    document.add_paragraph(
        "Synthetic test resume for CLI browser-approve/fill tests. No real content."
    )
    document.save(str(candidate_dir / "resume_master.docx"))
    return candidate_dir


def _configure_env(tmp_path: Path, monkeypatch, real_config) -> Path:
    cfg_dir = tmp_path / "config"
    shutil.copytree(real_config.env.config_dir, cfg_dir)
    candidate_dir = _make_synthetic_candidate_dir(tmp_path / "candidate")
    db_path = tmp_path / "cli_browser_approve_fill_test.db"
    monkeypatch.setenv("CONFIG_DIR", str(cfg_dir))
    monkeypatch.setenv("CANDIDATE_DIR", str(candidate_dir))
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{db_path}")
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    monkeypatch.setenv("COLUMNS", "250")
    return db_path


def _seed_application(
    db_path: Path, profile: CandidateProfile, *, application_url: str, fingerprint: str
) -> tuple[int, int]:
    engine = get_engine(f"sqlite:///{db_path}")
    init_db(engine)
    factory = get_session_factory(engine)
    with factory() as session:
        candidate_id = save_candidate_profile(session, profile)

        company = Company(name=f"Acme-{fingerprint}")
        session.add(company)
        session.flush()
        source = JobSource(
            name=f"cli-browser-approve-fill-{fingerprint}", kind="ats_api", enabled=True
        )
        session.add(source)
        session.flush()
        job = JobRow(
            source_id=source.id, source_job_id=fingerprint, company_id=company.id,
            company_name="Drivetrain", title="Business Analyst - Customer Platform",
            application_url=application_url, job_fingerprint=fingerprint,
        )
        session.add(job)
        session.flush()

        application = Application(job_id=job.id, candidate_id=candidate_id, status="MATCHED")
        session.add(application)
        session.commit()
        return application.id, job.id


class _NonClosingBrowserProxy:
    def __init__(self, real_browser):
        self._real_browser = real_browser

    def __getattr__(self, name):
        return getattr(self._real_browser, name)

    def close(self) -> None:
        pass


@pytest.fixture()
def patched_sync_playwright(monkeypatch, browser):
    calls: list[dict] = []

    class _FakeChromium:
        @staticmethod
        def launch(**kwargs):
            calls.append(dict(kwargs))
            return _NonClosingBrowserProxy(browser)

    class _FakePlaywright:
        chromium = _FakeChromium()

        def __enter__(self):
            return self

        def __exit__(self, *exc_info):
            return False

    monkeypatch.setattr(playwright_sync_api, "sync_playwright", lambda: _FakePlaywright())
    return calls


def _drivetrain_native_url(site_server: str) -> str:
    return f"{site_server}/drivetrain_business_analyst_native.html"


def _approve_args(application_id: int, url: str, *, extra: list[str] | None = None) -> list[str]:
    args = [
        "applications", "browser-approve", str(application_id), "--url", url,
        "--confirm", "--yes", "--answer", f"Current company={SYNTHETIC_VALUES['current_company']}",
    ]
    return args + (extra or [])


def _fill_args(application_id: int, url: str, *, extra: list[str] | None = None) -> list[str]:
    args = [
        "applications", "browser-fill", str(application_id), "--url", url,
        "--confirm", "--answer", f"Current company={SYNTHETIC_VALUES['current_company']}",
    ]
    return args + (extra or [])


# ==========================================================================
# Safe-by-default: no --confirm makes no network call, for both commands.
# ==========================================================================
def test_browser_approve_without_confirm_makes_no_network_call(
    tmp_path, monkeypatch, real_config, patched_sync_playwright
):
    _configure_env(tmp_path, monkeypatch, real_config)
    result = runner.invoke(
        app, ["applications", "browser-approve", "999", "--url", "https://example.test/apply"]
    )
    assert result.exit_code == 0
    assert "No action taken" in result.output
    assert patched_sync_playwright == []


def test_browser_fill_without_confirm_makes_no_network_call(
    tmp_path, monkeypatch, real_config, patched_sync_playwright
):
    _configure_env(tmp_path, monkeypatch, real_config)
    result = runner.invoke(
        app, ["applications", "browser-fill", "999", "--url", "https://example.test/apply"]
    )
    assert result.exit_code == 0
    assert "No action taken" in result.output
    assert patched_sync_playwright == []


# ==========================================================================
# UNPREPARED -> PREPARED -> HUMAN_REVIEW -> APPROVED -> FILL -> FILLED -> STOP
# ==========================================================================
def test_full_state_transition_approve_then_fill(
    tmp_path, monkeypatch, real_config, patched_sync_playwright, site_server
):
    db_path = _configure_env(tmp_path, monkeypatch, real_config)
    url = _drivetrain_native_url(site_server)
    application_id, job_id = _seed_application(
        db_path, _synthetic_identity_profile(), application_url=url,
        fingerprint="full-state-transition",
    )

    # UNPREPARED -> PREPARED -> HUMAN_REVIEW -> APPROVED
    approve_result = runner.invoke(app, _approve_args(application_id, url))
    assert approve_result.exit_code == 0, approve_result.output
    assert "Approved." in approve_result.output
    assert SYNTHETIC_VALUES["current_company"] in approve_result.output
    assert "Human input" in approve_result.output
    for value in (
        SYNTHETIC_VALUES["full_name"], SYNTHETIC_VALUES["email"], SYNTHETIC_VALUES["phone"],
    ):
        assert value in approve_result.output

    engine = get_engine(f"sqlite:///{db_path}")
    with get_session_factory(engine)() as session:
        approvals = session.execute(
            select(ApplicationApproval).where(
                ApplicationApproval.application_id == application_id
            )
        ).scalars().all()
        assert len(approvals) == 1
        assert approvals[0].consumed_at is None

    # FILL -> FILLED -> STOP
    fill_result = runner.invoke(app, _fill_args(application_id, url))
    assert fill_result.exit_code == 0, fill_result.output
    assert "STOPPED BEFORE SUBMISSION" in fill_result.output
    assert SYNTHETIC_VALUES["current_company"] in fill_result.output

    with get_session_factory(engine)() as session:
        approval = session.get(ApplicationApproval, approvals[0].id)
        assert approval.consumed_at is not None  # single-use, now spent

        events = session.execute(
            select(ApplicationEvent).where(
                ApplicationEvent.application_id == application_id,
                ApplicationEvent.event_type == "BROWSER_FILL_COMPLETED",
            )
        ).scalars().all()
        assert len(events) == 1

        # STOP: status untouched by this pipeline, exactly like browser-preview.
        application = session.get(Application, application_id)
        assert application.status == "MATCHED"


def test_fill_without_any_prior_approval_refuses(
    tmp_path, monkeypatch, real_config, patched_sync_playwright, site_server
):
    db_path = _configure_env(tmp_path, monkeypatch, real_config)
    url = _drivetrain_native_url(site_server)
    application_id, _ = _seed_application(
        db_path, _synthetic_identity_profile(), application_url=url, fingerprint="no-approval",
    )

    result = runner.invoke(app, _fill_args(application_id, url))
    assert result.exit_code == 1, result.output
    assert "No valid approval matches" in result.output
    assert "STOPPED BEFORE SUBMISSION" not in result.output


def test_approval_is_single_use_second_fill_refuses(
    tmp_path, monkeypatch, real_config, patched_sync_playwright, site_server
):
    db_path = _configure_env(tmp_path, monkeypatch, real_config)
    url = _drivetrain_native_url(site_server)
    application_id, _ = _seed_application(
        db_path, _synthetic_identity_profile(), application_url=url, fingerprint="single-use",
    )

    assert runner.invoke(app, _approve_args(application_id, url)).exit_code == 0
    first_fill = runner.invoke(app, _fill_args(application_id, url))
    assert first_fill.exit_code == 0, first_fill.output

    second_fill = runner.invoke(app, _fill_args(application_id, url))
    assert second_fill.exit_code == 1, second_fill.output
    assert "No valid approval matches" in second_fill.output


def test_fill_with_different_answer_than_what_was_approved_refuses(
    tmp_path, monkeypatch, real_config, patched_sync_playwright, site_server
):
    """The approval is bound to the EXACT prepared state -- a different
    --answer at fill time changes the snapshot fingerprint, so the stored
    approval no longer matches. Must refuse, never silently proceed with
    the different value."""
    db_path = _configure_env(tmp_path, monkeypatch, real_config)
    url = _drivetrain_native_url(site_server)
    application_id, _ = _seed_application(
        db_path, _synthetic_identity_profile(), application_url=url, fingerprint="drift",
    )

    assert runner.invoke(app, _approve_args(application_id, url)).exit_code == 0

    drifted = runner.invoke(
        app,
        [
            "applications", "browser-fill", str(application_id), "--url", url, "--confirm",
            "--answer", "Current company=A Totally Different Employer",
        ],
    )
    assert drifted.exit_code == 1, drifted.output
    assert "No valid approval matches" in drifted.output
    assert "A Totally Different Employer" not in drifted.output


def test_approve_refuses_when_a_question_remains_unresolved(
    tmp_path, monkeypatch, real_config, patched_sync_playwright, site_server
):
    """Mirrors the existing `applications approve` command's own
    precondition (reject if any answer still requires_human) -- an
    approval must never be created over an incomplete answer set."""
    db_path = _configure_env(tmp_path, monkeypatch, real_config)
    url = _drivetrain_native_url(site_server)
    application_id, _ = _seed_application(
        db_path, _synthetic_identity_profile(), application_url=url, fingerprint="incomplete",
    )

    result = runner.invoke(
        app,
        [
            "applications", "browser-approve", str(application_id),
            "--url", url, "--confirm", "--yes",
        ],
    )
    assert result.exit_code == 1, result.output
    assert "still has question(s) requiring human input" in result.output
    assert "Current company" in result.output

    engine = get_engine(f"sqlite:///{db_path}")
    with get_session_factory(engine)() as session:
        approvals = session.execute(
            select(ApplicationApproval).where(
                ApplicationApproval.application_id == application_id
            )
        ).scalars().all()
        assert approvals == []


def test_fill_stops_on_a_freshly_appeared_captcha(
    tmp_path, monkeypatch, real_config, patched_sync_playwright, site_server
):
    """Step 9: if the freshly re-inspected form now shows a security
    condition that wasn't there at approval time, browser-fill must stop
    before typing anything -- never silently proceed with a stale
    approval against a now-different, unreviewed form."""
    db_path = _configure_env(tmp_path, monkeypatch, real_config)
    approve_url = _drivetrain_native_url(site_server)
    application_id, job_id = _seed_application(
        db_path, _synthetic_identity_profile(), application_url=approve_url,
        fingerprint="captcha-after-approval",
    )
    assert runner.invoke(app, _approve_args(application_id, approve_url)).exit_code == 0

    # Simulate the target now showing a CAPTCHA -- point the SAME job at a
    # captcha fixture and update application_url accordingly (equivalent
    # to the real target's DOM having changed since approval).
    captcha_url = f"{site_server}/index.html?scenario=captcha"
    engine = get_engine(f"sqlite:///{db_path}")
    with get_session_factory(engine)() as session:
        job = session.get(JobRow, job_id)
        job.application_url = captcha_url
        session.commit()

    result = runner.invoke(app, _fill_args(application_id, captcha_url))
    assert result.exit_code == 0, result.output
    assert "Human review required" in result.output
    assert "nothing was filled" in result.output
    assert "STOPPED BEFORE SUBMISSION" not in result.output


# ==========================================================================
# Resume safety (Step 6): merely having a resume available must never
# cause an upload, for either new command.
# ==========================================================================
def test_browser_approve_never_uploads_resume(
    tmp_path, monkeypatch, real_config, patched_sync_playwright, site_server
):
    db_path = _configure_env(tmp_path, monkeypatch, real_config)
    url = _drivetrain_native_url(site_server)
    application_id, _ = _seed_application(
        db_path, _synthetic_identity_profile(), application_url=url, fingerprint="resume-approve",
    )
    result = runner.invoke(app, _approve_args(application_id, url))
    assert result.exit_code == 0, result.output
    assert "Uploaded files" not in result.output
    assert "Optional, not attached" in result.output


def test_browser_fill_never_uploads_resume(
    tmp_path, monkeypatch, real_config, patched_sync_playwright, site_server
):
    db_path = _configure_env(tmp_path, monkeypatch, real_config)
    url = _drivetrain_native_url(site_server)
    application_id, _ = _seed_application(
        db_path, _synthetic_identity_profile(), application_url=url, fingerprint="resume-fill",
    )
    assert runner.invoke(app, _approve_args(application_id, url)).exit_code == 0
    result = runner.invoke(app, _fill_args(application_id, url))
    assert result.exit_code == 0, result.output
    assert "Uploaded files" not in result.output
    assert "Optional, not attached" in result.output


# ==========================================================================
# Security audit: neither command ever calls submit()/verify() or touches
# allowlist code -- static check over each function body, mirroring
# test_human_input.py's forbidden-call-pattern convention.
# ==========================================================================
def test_browser_approve_and_fill_never_call_submit_verify_or_allowlist_code():
    source = Path("src/job_agent/cli/main.py").read_text()
    start = source.index('@applications_app.command("browser-approve")')
    end = source.index('@applications_app.command("run")')
    body = source[start:end]
    forbidden = (
        "create_allowlist_entry(", "revoke_allowlist_entry(",
        "provider.submit(", "provider.verify(",
        "submit_application(", "verify_application(",
        "EnvCredentialStore", "SubmissionHttpClient",
    )
    for term in forbidden:
        assert term not in body
