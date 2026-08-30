"""Phase 6D Stage 2 — CLI integration tests for `applications browser-preview`
(`typer.testing.CliRunner`, in-process, against a temporary sqlite database
and a temporary copy of the real config directory — the same convention
`test_cli_applications_phase6c.py` already established).

Every test that reaches the live-network branch (`--confirm` passed and
the `--url` matches) points at a LOCAL synthetic fixture served on
127.0.0.1 by the shared `site_server` fixture — never a real site.

PRIVACY NOTE (Phase 6D trusted-fact checkpoint): `applications browser-
preview` calls `parse_candidate_profile(cfg)` and reads `resume_master.
docx` from `cfg.env.candidate_dir` INDEPENDENTLY of whatever profile is
seeded into the test database — `cfg.env.candidate_dir` defaults to the
repository's real `candidate/` directory unless the `CANDIDATE_DIR` env
var is overridden. Since PR #11 made `generate_answer()` resolve verified
identity/contact facts deterministically, any test that left this
pointed at the real directory would have the REAL candidate's real name/
email/phone/location/LinkedIn URL typed into the local fixture's DOM and
rendered into `result.output` — and, because the CLI always passes
`resume_path` to `BrowserApplicationProvider` unconditionally, the real
`resume_master.docx` would be read (and, once a fill reaches a file
field, attached) too. `_configure_env()` below builds and points
`CANDIDATE_DIR` at a MINIMAL, entirely synthetic candidate directory
(`_make_synthetic_candidate_dir()`) for every test in this file, so
`result.output` — printed verbatim by several assertions below on
failure — can only ever contain the same synthetic placeholders already
established elsewhere in this test suite.

The CLI command owns its own Playwright lifecycle end to end (launch,
use, close) — the right design for a real, standalone CLI invocation.
That conflicts with this test session's shared, session-scoped `browser`
fixture (tests/conftest.py): Playwright's sync API allows only one
active `sync_playwright()` context per thread, and pytest runs a whole
session in one thread. `patched_sync_playwright` below redirects the
CLI's own `sync_playwright()`/`chromium.launch()` calls to reuse that
already-open shared browser instead of opening a second one, and makes
the CLI's own `browser.close()` a no-op so it never tears down the
shared browser out from under every other test still using it. This
changes nothing about the CLI command's real, production behavior —
only how this test suite avoids a same-process resource conflict that
only exists because many tests share one process.
"""

from __future__ import annotations

import shutil
from pathlib import Path

import playwright.sync_api as playwright_sync_api
import pytest
from typer.testing import CliRunner

from job_agent.candidate.schema import (
    CandidateProfile,
    Fact,
    LocationPreferences,
    SalaryPreferences,
    TargetRoles,
    VisaInformation,
    WorkPreferences,
)
from job_agent.cli.main import app
from job_agent.db.models import Application, Company, JobSource
from job_agent.db.models import Job as JobRow
from job_agent.db.repository import save_candidate_profile
from job_agent.db.session import get_engine, get_session_factory, init_db

pytest.importorskip("playwright.sync_api")

CHROMIUM_EXECUTABLE = "/opt/pw-browsers/chromium-1194/chrome-linux/chrome"

pytestmark = pytest.mark.skipif(
    not Path(CHROMIUM_EXECUTABLE).exists(),
    reason="pre-installed Chromium not available in this environment",
)

runner = CliRunner()

# Obvious synthetic placeholders only — the exact same values already
# established in tests/unit/test_drivetrain_target_compatibility.py —
# never real personal data.
SYNTHETIC_VALUES = {
    "full_name": "Test Candidate",
    "email": "test.candidate@example.invalid",
    "phone": "+1-555-0100",
    "current_location": "Test City, Test Country",
    "linkedin_url": "https://linkedin.com/in/testcandidate",
}


def _fact(value: str) -> Fact[str]:
    return Fact[str](
        value=value, source="test_synthetic_profile", confidence=1.0, verified=True
    )


def _synthetic_identity_profile() -> CandidateProfile:
    """A fully synthetic CandidateProfile — mirrors the identical helper
    in test_drivetrain_target_compatibility.py. Used to seed the test
    database; the CLI's OWN answer generation re-parses a profile
    straight from `cfg.env.candidate_dir` independently of this, which is
    why `_make_synthetic_candidate_dir()` below matters just as much."""
    unknown_pref = Fact.unknown(source="test_synthetic_profile")
    return CandidateProfile(
        identity_name=_fact(SYNTHETIC_VALUES["full_name"]),
        identity_current_location=_fact(SYNTHETIC_VALUES["current_location"]),
        contact_email=_fact(SYNTHETIC_VALUES["email"]),
        contact_phone=_fact(SYNTHETIC_VALUES["phone"]),
        contact_linkedin=_fact(SYNTHETIC_VALUES["linkedin_url"]),
        target_roles=TargetRoles(),
        work_preferences=WorkPreferences(
            remote=unknown_pref, willing_to_relocate=unknown_pref, notice_period=unknown_pref
        ),
        location_preferences=LocationPreferences(
            current_location=_fact(SYNTHETIC_VALUES["current_location"]),
            open_to_countries=unknown_pref,
        ),
        salary_preferences=SalaryPreferences(
            currency=unknown_pref, minimum_annual=unknown_pref,
            target_annual=unknown_pref, negotiable=unknown_pref,
        ),
        visa_information=VisaInformation(
            nationality=unknown_pref, requires_sponsorship_us=unknown_pref,
            requires_sponsorship_uk=unknown_pref, requires_sponsorship_eu=unknown_pref,
            requires_sponsorship_other=unknown_pref,
        ),
    )


def _make_synthetic_candidate_dir(candidate_dir: Path) -> Path:
    """Builds the MINIMAL, entirely synthetic candidate/ directory tree
    `parse_candidate_profile()`/`extract_resume_text()` need to succeed —
    see the module docstring's PRIVACY NOTE for why this exists. Every
    per-entity file (experience/projects/skills/education/achievements)
    is deliberately empty: `_read_lines()` only raises if a file is
    MISSING, not if it has no content, so an empty file parses to zero
    entries without error."""
    candidate_dir.mkdir(parents=True, exist_ok=True)
    (candidate_dir / "profile.md").write_text(
        "## Identity\n"
        f"name: {SYNTHETIC_VALUES['full_name']}\n"
        f"current_location: {SYNTHETIC_VALUES['current_location']}\n"
        "\n"
        "## Contact\n"
        f"email: {SYNTHETIC_VALUES['email']}\n"
        f"phone: {SYNTHETIC_VALUES['phone']}\n"
        f"linkedin: {SYNTHETIC_VALUES['linkedin_url']}\n"
    )
    for filename in (
        "experience.md", "projects.md", "skills.md", "education.md", "achievements.md",
    ):
        (candidate_dir / filename).write_text("")
    (candidate_dir / "answers").mkdir(exist_ok=True)

    import docx

    document = docx.Document()
    document.add_paragraph(
        "Synthetic test resume for CLI browser-preview tests. No real content."
    )
    document.save(str(candidate_dir / "resume_master.docx"))
    return candidate_dir


def _configure_env(tmp_path: Path, monkeypatch, real_config) -> Path:
    cfg_dir = tmp_path / "config"
    shutil.copytree(real_config.env.config_dir, cfg_dir)
    candidate_dir = _make_synthetic_candidate_dir(tmp_path / "candidate")
    db_path = tmp_path / "cli_browser_preview_test.db"
    monkeypatch.setenv("CONFIG_DIR", str(cfg_dir))
    monkeypatch.setenv("CANDIDATE_DIR", str(candidate_dir))
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{db_path}")
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    monkeypatch.setenv("COLUMNS", "250")
    return db_path


def _seed_application(
    db_path: Path, profile: CandidateProfile, *, application_url: str, fingerprint: str
) -> tuple[int, int]:
    engine = get_engine(f"sqlite:///{db_path}")
    init_db(engine)
    factory = get_session_factory(engine)
    with factory() as session:
        candidate_id = save_candidate_profile(session, profile)

        company = Company(name=f"Acme-{fingerprint}")
        session.add(company)
        session.flush()
        source = JobSource(name=f"cli-browser-preview-{fingerprint}", kind="ats_api", enabled=True)
        session.add(source)
        session.flush()
        job = JobRow(
            source_id=source.id, source_job_id=fingerprint, company_id=company.id,
            company_name="Drivetrain", title="Business Analyst - Customer Platform",
            application_url=application_url, job_fingerprint=fingerprint,
        )
        session.add(job)
        session.flush()

        application = Application(job_id=job.id, candidate_id=candidate_id, status="MATCHED")
        session.add(application)
        session.commit()
        return application.id, job.id


class _NonClosingBrowserProxy:
    """Delegates everything to the wrapped Browser except close(), which
    is a no-op. The CLI command calls browser.close() itself when it's
    done — entirely correct for a real, standalone invocation — but here
    the "browser" it's holding is this test session's shared,
    session-scoped fixture; closing it for real would break every other
    test still relying on it for the rest of the session."""

    def __init__(self, real_browser):
        self._real_browser = real_browser

    def __getattr__(self, name):
        return getattr(self._real_browser, name)

    def close(self) -> None:
        pass


@pytest.fixture()
def patched_sync_playwright(monkeypatch, browser):
    """Redirects the CLI command's own `sync_playwright()`/
    `chromium.launch()` calls to reuse the already-open, shared `browser`
    fixture instead of opening a second Playwright context — see module
    docstring for why. Returns the list of "launch" calls so a test can
    assert whether the CLI ever reached that point at all."""
    calls: list[dict] = []

    class _FakeChromium:
        @staticmethod
        def launch(**kwargs):
            calls.append(dict(kwargs))
            return _NonClosingBrowserProxy(browser)

    class _FakePlaywright:
        chromium = _FakeChromium()

        def __enter__(self):
            return self

        def __exit__(self, *exc_info):
            return False

    monkeypatch.setattr(playwright_sync_api, "sync_playwright", lambda: _FakePlaywright())
    return calls


def _drivetrain_native_url(site_server: str) -> str:
    return f"{site_server}/drivetrain_business_analyst_native.html"


# ==========================================================================
# Safe-by-default: no --confirm makes no network call.
# ==========================================================================
def test_without_confirm_makes_no_network_call(
    tmp_path, monkeypatch, real_config, patched_sync_playwright
):
    _configure_env(tmp_path, monkeypatch, real_config)
    result = runner.invoke(
        app, ["applications", "browser-preview", "999", "--url", "https://example.test/apply"]
    )
    assert result.exit_code == 0
    assert "No action taken" in result.output
    assert patched_sync_playwright == []  # never launched


def test_unknown_application_id_errors_before_any_network_call(
    tmp_path, monkeypatch, real_config, patched_sync_playwright
):
    _configure_env(tmp_path, monkeypatch, real_config)
    result = runner.invoke(
        app,
        [
            "applications", "browser-preview", "999999",
            "--url", "https://example.test/apply", "--confirm",
        ],
    )
    assert result.exit_code == 1
    assert "No application with id" in result.output
    assert patched_sync_playwright == []


def test_url_mismatch_rejected_before_any_network_call(
    tmp_path, monkeypatch, real_config, patched_sync_playwright
):
    db_path = _configure_env(tmp_path, monkeypatch, real_config)
    application_id, _ = _seed_application(
        db_path, _synthetic_identity_profile(),
        application_url="https://real.example.test/apply", fingerprint="url-mismatch",
    )
    result = runner.invoke(
        app,
        [
            "applications", "browser-preview", str(application_id),
            "--url", "https://a-different-url.example.test/apply", "--confirm",
        ],
    )
    assert result.exit_code == 1
    assert "does not match this job's own application_url" in result.output
    assert patched_sync_playwright == []


# ==========================================================================
# Live path — always against a local synthetic fixture only.
# ==========================================================================
def test_happy_path_renders_snapshot_five_trusted_fields_resolved_current_company_needs_input(
    tmp_path, monkeypatch, real_config, patched_sync_playwright, site_server
):
    """Phase 6D trusted-fact checkpoint: since PR #11, `generate_answer()`
    resolves "Full name"/"Email"/"Phone"/"Current location"/"LinkedIn URL"
    deterministically from CandidateProfile's verified Fact[str] fields —
    never from the LLM (no ANTHROPIC_API_KEY is configured; the autouse
    fixture strips it and this test never sets one) and never from the
    answer bank (empty in the synthetic candidate dir). "Current company"
    has no trusted-fact mapping and correctly remains HUMAN_REQUIRED. The
    synthetic candidate dir built by `_configure_env()` means every value
    that could appear in `result.output` — printed verbatim on assertion
    failure — is one of the SYNTHETIC_VALUES placeholders above, never
    real personal data."""
    db_path = _configure_env(tmp_path, monkeypatch, real_config)
    url = _drivetrain_native_url(site_server)
    application_id, _ = _seed_application(
        db_path, _synthetic_identity_profile(), application_url=url, fingerprint="happy-path",
    )

    result = runner.invoke(
        app, ["applications", "browser-preview", str(application_id), "--url", url, "--confirm"]
    )

    assert result.exit_code == 0, result.output
    assert len(patched_sync_playwright) == 1
    assert "Drivetrain" in result.output
    assert "Business Analyst - Customer Platform" in result.output
    assert "Needs your input" in result.output

    # The five trusted fields resolved: their synthetic values are shown
    # as PROPOSED (filled) values, never fabricated, never real data.
    for synthetic_value in SYNTHETIC_VALUES.values():
        assert synthetic_value in result.output

    # "Current company" (no trusted-fact mapping) and the resume (never
    # attached — see PRIVACY NOTE) both still need a human, and neither
    # field id silently disappears from the rendered output.
    assert "current_company" in result.output
    assert "structurally unavailable" in result.output
    # Real-target-readiness checkpoint: the resume must never be
    # auto-attached just because resume_master.docx happens to exist in
    # the candidate dir -- it is disclosed as optional and unattached,
    # not silently uploaded.
    assert "Uploaded files" not in result.output
    assert "Optional, not attached" in result.output
    # Real-target-readiness checkpoint: unresolved "Current company" is
    # disclosed WITH a reason, not just a bare unexplained "?".
    assert "Human decision required" in result.output


# ==========================================================================
# Real-target-readiness checkpoint: --answer lets a human resolve a
# CURRENTLY unresolved question (e.g. "Current company") at review time,
# without ever touching CandidateProfile or the LLM, and always
# distinguishable in the output from a trusted candidate fact.
# ==========================================================================
def test_answer_override_resolves_current_company_and_is_labeled_human_input(
    tmp_path, monkeypatch, real_config, patched_sync_playwright, site_server
):
    db_path = _configure_env(tmp_path, monkeypatch, real_config)
    url = _drivetrain_native_url(site_server)
    application_id, _ = _seed_application(
        db_path, _synthetic_identity_profile(), application_url=url,
        fingerprint="human-input-override",
    )

    result = runner.invoke(
        app,
        [
            "applications", "browser-preview", str(application_id),
            "--url", url, "--confirm", "--answer", "Current company=Instawork",
        ],
    )

    assert result.exit_code == 0, result.output
    assert "Instawork" in result.output
    assert "Human input" in result.output
    # The five trusted fields are STILL trusted facts -- an unrelated
    # override never changes how they resolved.
    for synthetic_value in SYNTHETIC_VALUES.values():
        assert synthetic_value in result.output
    assert result.output.count("Trusted candidate fact") == 5
    # "Current company" is now fully resolved -- with it filled in,
    # nothing on this form remains unresolved at all.
    assert "Needs your input" not in result.output


def test_without_answer_override_current_company_still_needs_human_input(
    tmp_path, monkeypatch, real_config, patched_sync_playwright, site_server
):
    """Baseline: omitting --answer entirely leaves "Current company"
    exactly as unresolved as before -- the override mechanism is opt-in,
    never a default behavior change."""
    db_path = _configure_env(tmp_path, monkeypatch, real_config)
    url = _drivetrain_native_url(site_server)
    application_id, _ = _seed_application(
        db_path, _synthetic_identity_profile(), application_url=url,
        fingerprint="no-override",
    )

    result = runner.invoke(
        app, ["applications", "browser-preview", str(application_id), "--url", url, "--confirm"]
    )

    assert result.exit_code == 0, result.output
    assert "Instawork" not in result.output
    assert "Human input" not in result.output
    assert "Human decision required" in result.output


def test_answer_override_never_modifies_an_already_resolved_trusted_fact(
    tmp_path, monkeypatch, real_config, patched_sync_playwright, site_server
):
    """A human cannot use --answer to silently override a field the
    system already resolved from a trusted CandidateProfile fact."""
    db_path = _configure_env(tmp_path, monkeypatch, real_config)
    url = _drivetrain_native_url(site_server)
    application_id, _ = _seed_application(
        db_path, _synthetic_identity_profile(), application_url=url,
        fingerprint="override-rejected",
    )

    result = runner.invoke(
        app,
        [
            "applications", "browser-preview", str(application_id),
            "--url", url, "--confirm", "--answer", "Email=attacker@example.test",
        ],
    )

    assert result.exit_code == 0, result.output
    assert "attacker@example.test" not in result.output
    assert SYNTHETIC_VALUES["email"] in result.output
    assert "was not applied" in result.output


def test_malformed_answer_option_rejected_before_any_network_call(
    tmp_path, monkeypatch, real_config, patched_sync_playwright
):
    _configure_env(tmp_path, monkeypatch, real_config)
    result = runner.invoke(
        app,
        [
            "applications", "browser-preview", "1",
            "--url", "https://example.test/apply", "--confirm",
            "--answer", "no-equals-sign-here",
        ],
    )
    assert result.exit_code == 1
    assert patched_sync_playwright == []


def test_captcha_scenario_stops_before_filling_anything(
    tmp_path, monkeypatch, real_config, patched_sync_playwright, site_server
):
    db_path = _configure_env(tmp_path, monkeypatch, real_config)
    url = f"{site_server}/index.html?scenario=captcha"
    application_id, _ = _seed_application(
        db_path, _synthetic_identity_profile(), application_url=url,
        fingerprint="captcha-scenario",
    )

    result = runner.invoke(
        app, ["applications", "browser-preview", str(application_id), "--url", url, "--confirm"]
    )

    assert result.exit_code == 0, result.output
    assert "Human review required" in result.output
    assert "nothing was filled" in result.output
    assert "Needs your input" not in result.output  # never reached the snapshot stage
