"""Deterministic raw-text extraction from the authoritative resume file.

No LLM involved — this is pure structural text extraction (python-docx),
consistent with this codebase's preference for deterministic code over
LLM calls wherever one suffices (BUILD PROMPT section 59). It exists so
`job_agent.resume.validator` has ground truth to check the parsed
CandidateProfile against.

Only `.docx` is supported today because that's what `candidate/
resume_master.docx` actually is. A `.pdf` resume would need a different
extraction path (BUILD PROMPT section 55 lists broader format support as
a future extension point, not part of Phase 4's scope).
"""

from __future__ import annotations

from pathlib import Path

from job_agent.resume.errors import ResumeExtractionError


def extract_resume_text(path: Path) -> str:
    """Extract all paragraph and table text from a .docx resume.

    Raises `ResumeExtractionError` on any failure — missing file, wrong
    format, corrupt file, or a file with no extractable text — rather than
    returning an empty/partial string that downstream validation could
    mistake for "nothing to check against".
    """
    if not path.exists():
        raise ResumeExtractionError(f"Resume file not found: {path}")

    try:
        import docx
    except ImportError as exc:  # pragma: no cover - dependency is declared, always installed
        raise ResumeExtractionError(
            "python-docx is not installed; cannot read the resume file"
        ) from exc

    try:
        document = docx.Document(str(path))
    except Exception as exc:  # noqa: BLE001 - any file-format/corruption failure
        raise ResumeExtractionError(f"Could not open resume file {path}: {exc}") from exc

    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)

    text = "\n".join(parts)
    if not text.strip():
        raise ResumeExtractionError(f"Resume file {path} contained no extractable text")
    return text
